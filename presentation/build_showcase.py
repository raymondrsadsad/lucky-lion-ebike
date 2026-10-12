#!/usr/bin/env python3
"""Build a 'desktop + mobile of every part' showcase deck.

Run the dev server first:  python3 server.py 8000
Seed the demo data first:  supabase/demo-reset.sql (already applied live)
Then:                      python3 presentation/build_showcase.py

For each part of the system it captures a DESKTOP shot and a MOBILE shot of the
CURRENT (redesigned) UI, then assembles one slide per part (desktop left, phone
right) into presentation/Lucky-Lion-Showcase.pptx.

Standalone — does not touch the main capture.py / build_deck.py.
"""
import os, sys, traceback
from playwright.sync_api import sync_playwright
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

BASE = "http://localhost:8000"
HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "screenshots", "showcase")
os.makedirs(OUT, exist_ok=True)

DESKTOP = {"width": 1440, "height": 900}
MOBILE = {"width": 390, "height": 844}
CUST, STAFF, ADMIN = ("user", "user123"), ("staff", "staff123"), ("admin", "admin123")

RED = RGBColor(0xE1, 0x1B, 0x22)
INK = RGBColor(0x14, 0x14, 0x14)
MUT = RGBColor(0x6B, 0x70, 0x76)
PAPER = RGBColor(0xF4, 0xF5, 0xF7)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

# Order of the deck. Each part → one slide (desktop + mobile).
PARTS = [
    ("home",         "Homepage",             "Customer landing — hero, branches and the Products entry."),
    ("categories",   "Shop by wheel type",   "Homepage category tiles — 4 / 3 / 2 wheels into the full lineup."),
    ("products",     "Products catalog",     "Full lineup, filterable by wheel type (4 / 3 / 2 wheels)."),
    ("product",      "Product detail",       "A dedicated page per unit, with its price and specs."),
    ("location",     "Branch locations",     "All 9 Cavite branches with call + get-directions."),
    ("about",        "About the company",    "Company story and the World-Class Philippines award."),
    ("menu",         "Account menu",         "The responsive icon-tile burger menu (the new look)."),
    ("login",        "Sign in",              "Customer / staff / admin log in from one place."),
    ("account",      "Customer dashboard",   "A member's home: their e-bikes, bookings and messages."),
    ("schedule",     "Schedule a repair",    "Customers book a service slot for a unit they own."),
    ("messages",     "Customer messages",    "Direct chat thread between a customer and the branch."),
    ("staffdash",    "Staff dashboard",      "Today's repairs, overdue, completed and unread inbox."),
    ("portalmenu",   "Portal menu",          "The staff/admin side-drawer as an icon-tile grid."),
    ("repairs",      "Scheduled repairs",    "Manage bookings: status, cost, edit, complete, history."),
    ("registry",     "E-Bike registry",      "Every unit sold across customers, searchable by serial."),
    ("users",        "Customers & units",    "Per-branch customer list with their owned e-bikes."),
    ("models",       "Models manager",       "Edit specs, colours and photos for each model."),
    ("poster",       "Poster maker",         "Generate a branded sale flyer, live, then download it."),
    ("settings",     "Site settings",        "Per-branch contacts, maps, socials and the sale theme."),
    ("analytics",    "Admin analytics",      "Company-wide repairs by branch and leaderboards."),
]

log = []

# ------------------------------------------------------------------ helpers
def settle(page, ms=1200): page.wait_for_timeout(ms)

def kill_motion(page):
    try:
        page.add_style_tag(content="*,*::before,*::after{animation:none!important;transition:none!important}")
    except Exception: pass

def snap(page, key, device):
    path = os.path.join(OUT, f"{key}.{device}.png")
    page.screenshot(path=path)
    log.append(f"{key}.{device}")
    print("  ✓", key, device)
    return path

def login(page, creds):
    # #loginBtn lives inside the guest burger dropdown now → open it first.
    page.click("#guestMenuBtn"); settle(page, 500)
    page.click("#loginBtn"); settle(page, 600)
    page.fill("#loginUser", creds[0])
    page.fill("#loginPass", creds[1])
    page.click("#loginForm button[type=submit]")
    settle(page, 2500)

def open_guest_menu(page):
    if page.eval_on_selector("#guestDropdown", "el => el.classList.contains('hidden')"):
        page.click("#guestMenuBtn"); settle(page, 500)

def open_member_menu(page):
    if page.eval_on_selector("#memberDropdown", "el => el.classList.contains('hidden')"):
        page.click("#hamburgerBtn"); settle(page, 500)

def open_panel(page, kind):
    open_member_menu(page)
    page.click(f"#memberDropdown button[data-panel={kind}]")
    settle(page, 1800)

def esc(page):
    page.keyboard.press("Escape"); settle(page, 400)

def hashto(page, h, wait=1600):
    page.evaluate("v => { location.hash = v; }", h)
    settle(page, wait)
    page.evaluate("window.scrollTo(0,0)"); settle(page, 300)

def guard(fn):
    try: fn()
    except Exception as e:
        print("  ! failed:", e); traceback.print_exc()

# ------------------------------------------------------------------ capture one device
def capture_device(browser, device, viewport, scale, is_mobile):
    print(f"== capturing {device} ==")
    # ---- public (logged out) ----
    ctx = browser.new_context(viewport=viewport, device_scale_factor=scale,
                              is_mobile=is_mobile, has_touch=is_mobile)
    page = ctx.new_page(); page.goto(BASE, wait_until="networkidle"); settle(page, 2500); kill_motion(page)
    guard(lambda: snap(page, "home", device))
    def categories_shot():
        page.evaluate("window.scrollTo(0,0)"); settle(page, 300)
        page.evaluate("document.querySelector('#store').scrollIntoView({block:'start',behavior:'instant'})")
        settle(page, 1200); snap(page, "categories", device)
        page.evaluate("window.scrollTo(0,0)"); settle(page, 300)
    guard(categories_shot)
    guard(lambda: (hashto(page, "#products", 1800), snap(page, "products", device)))
    guard(lambda: (hashto(page, "#product/lucky-agila-v2", 1600), snap(page, "product", device)))
    guard(lambda: (hashto(page, "#location", 1400), snap(page, "location", device)))
    guard(lambda: (hashto(page, "#about", 1400), snap(page, "about", device)))
    guard(lambda: (hashto(page, "#store", 900), open_guest_menu(page), settle(page, 500), snap(page, "menu", device)))
    def do_login_shot():
        # close the menu, reopen, click login → modal
        hashto(page, "#store", 500)
        open_guest_menu(page)
        page.click("#loginBtn"); settle(page, 700)
        snap(page, "login", device)
        esc(page)
    guard(do_login_shot)
    ctx.close()

    # ---- customer ----
    ctx = browser.new_context(viewport=viewport, device_scale_factor=scale,
                              is_mobile=is_mobile, has_touch=is_mobile)
    page = ctx.new_page(); page.goto(BASE, wait_until="networkidle"); settle(page, 2500); kill_motion(page)
    guard(lambda: login(page, CUST))
    guard(lambda: (hashto(page, "#account", 1600), snap(page, "account", device)))
    guard(lambda: (open_panel(page, "schedule"), snap(page, "schedule", device), esc(page)))
    guard(lambda: (open_panel(page, "messages"), settle(page, 1200), snap(page, "messages", device), esc(page)))
    ctx.close()

    # ---- staff ----
    ctx = browser.new_context(viewport=viewport, device_scale_factor=scale,
                              is_mobile=is_mobile, has_touch=is_mobile)
    page = ctx.new_page(); page.goto(BASE, wait_until="networkidle"); settle(page, 2500)
    guard(lambda: login(page, STAFF))
    page.goto(BASE + "/staff.html", wait_until="networkidle"); settle(page, 3000); kill_motion(page)
    if "staff.html" in page.url:
        guard(lambda: (hashto(page, "#dashboard", 2200), snap(page, "staffdash", device)))
        def portalmenu():
            page.click("#drawerBtn"); settle(page, 700)
            snap(page, "portalmenu", device)
            try: page.click("#drawerClose")
            except Exception: pass
            settle(page, 400)
        guard(portalmenu)
        guard(lambda: (hashto(page, "#repairs", 2000), snap(page, "repairs", device)))
        guard(lambda: (hashto(page, "#registry", 2000), snap(page, "registry", device)))
        guard(lambda: (hashto(page, "#users", 2000), snap(page, "users", device)))
        guard(lambda: (hashto(page, "#models", 2500), snap(page, "models", device)))
        guard(lambda: (hashto(page, "#poster", 2200), snap(page, "poster", device)))
        guard(lambda: (hashto(page, "#settings", 1600), snap(page, "settings", device)))
    else:
        print("  ! staff login redirected to", page.url)
    ctx.close()

    # ---- admin (analytics) ----
    ctx = browser.new_context(viewport=viewport, device_scale_factor=scale,
                              is_mobile=is_mobile, has_touch=is_mobile)
    page = ctx.new_page(); page.goto(BASE, wait_until="networkidle"); settle(page, 2500)
    guard(lambda: login(page, ADMIN))
    page.goto(BASE + "/staff.html", wait_until="networkidle"); settle(page, 3000); kill_motion(page)
    if "staff.html" in page.url:
        guard(lambda: (hashto(page, "#analytics", 2400), snap(page, "analytics", device)))
    else:
        print("  ! admin login redirected to", page.url)
    ctx.close()

def capture_all():
    with sync_playwright() as p:
        browser = p.chromium.launch()
        capture_device(browser, "desktop", DESKTOP, 2, False)
        capture_device(browser, "mobile", MOBILE, 3, True)
        browser.close()

# ------------------------------------------------------------------ deck
def px_size(path):
    from PIL import Image
    with Image.open(path) as im: return im.size

def build_deck():
    prs = Presentation()
    prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]
    SW, SH = prs.slide_width, prs.slide_height

    def rect(slide, x, y, w, h, fill, line=None, lw=None, round_=False):
        from pptx.enum.shapes import MSO_SHAPE
        shp = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE if round_ else MSO_SHAPE.RECTANGLE, x, y, w, h)
        shp.fill.solid(); shp.fill.fore_color.rgb = fill
        if line is not None:
            shp.line.color.rgb = line; shp.line.width = lw or Pt(1)
        else:
            shp.line.fill.background()
        shp.shadow.inherit = False
        return shp

    def textbox(slide, x, y, w, h, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
        tb = slide.shapes.add_textbox(x, y, w, h); tf = tb.text_frame
        tf.word_wrap = True; tf.vertical_anchor = anchor
        p = tf.paragraphs[0]; p.alignment = align
        for txt, size, color, bold in runs:
            r = p.add_run(); r.text = txt
            r.font.size = Pt(size); r.font.color.rgb = color; r.font.bold = bold
            r.font.name = "Poppins"
        return tb

    def fit(path, max_w, max_h):
        w, h = px_size(path); ar = w / h
        W = max_w; H = Emu(int(W / ar))
        if H > max_h: H = max_h; W = Emu(int(H * ar))
        return W, H

    # ---- title slide ----
    s = prs.slides.add_slide(blank)
    rect(s, 0, 0, SW, SH, INK)
    rect(s, 0, Inches(3.05), SW, Inches(0.05), RED)
    textbox(s, Inches(1), Inches(2.0), Inches(11.3), Inches(1.2),
            [("Lucky Lion eBike", 46, WHITE, True)], PP_ALIGN.CENTER)
    textbox(s, Inches(1), Inches(3.2), Inches(11.3), Inches(0.8),
            [("System features — desktop & mobile", 22, RGBColor(0xCB,0xD0,0xD6), False)], PP_ALIGN.CENTER)
    textbox(s, Inches(1), Inches(6.6), Inches(11.3), Inches(0.5),
            [("Responsive web system · Cavite · 9 branches", 13, MUT, False)], PP_ALIGN.CENTER)

    # ---- one slide per part ----
    made = 0
    for key, title, desc in PARTS:
        dpath = os.path.join(OUT, f"{key}.desktop.png")
        mpath = os.path.join(OUT, f"{key}.mobile.png")
        if not os.path.exists(dpath) and not os.path.exists(mpath):
            print("  · skip (no shots):", key); continue
        s = prs.slides.add_slide(blank)
        rect(s, 0, 0, SW, SH, WHITE)
        rect(s, 0, 0, Inches(0.28), SH, RED)                     # left accent
        textbox(s, Inches(0.6), Inches(0.35), Inches(12.2), Inches(0.7),
                [(title, 30, INK, True)])
        textbox(s, Inches(0.62), Inches(1.05), Inches(12.2), Inches(0.5),
                [(desc, 15, MUT, False)])

        area_top = Inches(1.75); area_bot = SH - Inches(0.55)
        area_h = area_bot - area_top

        # desktop (left, large)
        if os.path.exists(dpath):
            W, H = fit(dpath, Inches(8.3), area_h - Inches(0.45))
            x = Inches(0.6); y = area_top + Emu(int((area_h - H) / 2))
            rect(s, x - Emu(60000), y - Emu(60000), W + Emu(120000), H + Emu(120000), PAPER, INK, Pt(0.75))
            s.shapes.add_picture(dpath, x, y, width=W, height=H)
            textbox(s, x, y + H + Emu(30000), W, Inches(0.3),
                    [("Desktop", 12, MUT, True)], PP_ALIGN.CENTER)

        # mobile (right, phone)
        if os.path.exists(mpath):
            W, H = fit(mpath, Inches(2.7), area_h - Inches(0.45))
            x = Inches(10.2); y = area_top + Emu(int((area_h - H) / 2))
            rect(s, x - Emu(70000), y - Emu(70000), W + Emu(140000), H + Emu(140000), INK, None, None, round_=True)
            s.shapes.add_picture(mpath, x, y, width=W, height=H)
            textbox(s, x - Inches(0.4), y + H + Emu(40000), W + Inches(0.8), Inches(0.3),
                    [("Mobile", 12, MUT, True)], PP_ALIGN.CENTER)
        made += 1

    out = os.path.join(HERE, "Lucky-Lion-Showcase.pptx")
    prs.save(out)
    print(f"\nDeck: {made} feature slides → {out}")
    return out

def main():
    stage = set(sys.argv[1:]) or {"capture", "deck"}
    if "capture" in stage: capture_all()
    if "deck" in stage: build_deck()
    print(f"{len(log)} screenshots in {OUT}")

if __name__ == "__main__":
    main()
