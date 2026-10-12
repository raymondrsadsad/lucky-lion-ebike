#!/usr/bin/env python3
"""Assemble Lucky-Lion-Presentation.pptx from presentation/screenshots/.

  python3 presentation/build_deck.py

Re-run any time after capture.py to regenerate the deck.
"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

HERE = os.path.dirname(os.path.abspath(__file__))
SHOTS = os.path.join(HERE, "screenshots")
OUTFILE = os.path.join(HERE, "Lucky-Lion-Presentation.pptx")

RED = RGBColor(0xE1, 0x1B, 0x22)
INK = RGBColor(0x14, 0x14, 0x14)
MUTED = RGBColor(0x6B, 0x70, 0x7B)
LINE = RGBColor(0xE3, 0xE6, 0xEA)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BG = RGBColor(0xFA, 0xFA, 0xFB)

W, H = Inches(13.333), Inches(7.5)

prs = Presentation()
prs.slide_width, prs.slide_height = W, H
BLANK = prs.slide_layouts[6]


# ------------------------------------------------------------------ helpers
def add_slide(bg=BG):
    s = prs.slides.add_slide(BLANK)
    r = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, H)
    r.fill.solid()
    r.fill.fore_color.rgb = bg
    r.line.fill.background()
    r.shadow.inherit = False
    return s


def textbox(s, x, y, w, h, text, size=18, bold=False, color=INK,
            align=PP_ALIGN.LEFT, space_after=6, line=1.25):
    tb = s.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    lines = text if isinstance(text, list) else [text]
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = ln
        p.alignment = align
        p.space_after = Pt(space_after)
        p.line_spacing = line
        f = p.font
        f.size, f.bold, f.color.rgb, f.name = Pt(size), bold, color, "Calibri"
    return tb


def bullets(s, x, y, w, items, size=15):
    tb = s.shapes.add_textbox(x, y, w, Inches(4))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    hang = Inches(0.23)
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = "•  " + it
        p.space_after = Pt(11)
        p.line_spacing = 1.22
        p.font.size, p.font.color.rgb, p.font.name = Pt(size), INK, "Calibri"
        # hanging indent so wrapped lines line up with the text, not the dot
        pPr = p._p.get_or_add_pPr()
        pPr.set("marL", str(int(hang)))
        pPr.set("indent", str(int(-hang)))
    return tb


def est_bullet_height(items, w_in, size=15):
    """Rough height of a bullet block, for keeping badges clear of the text."""
    chars_per_line = max(18, int(w_in * 96 / (size * 0.52)))
    lines = sum(max(1, -(-len(it) // chars_per_line)) for it in items)
    return Inches(lines * size * 1.22 / 72 + len(items) * 11 / 72)


def picture(s, name, x, y, w=None, h=None, border=True):
    """Place a screenshot, sized by width OR height, keeping aspect ratio."""
    path = os.path.join(SHOTS, name + ".png")
    if not os.path.exists(path):
        print("  ! missing screenshot:", name)
        return None
    pic = s.shapes.add_picture(path, x, y, width=w, height=h)
    if border:
        pic.line.color.rgb = LINE
        pic.line.width = Pt(1)
    return pic


def crop(name, top=0.0, bottom=1.0):
    """Write a vertically cropped copy of a screenshot and return its stem.

    Used to focus a slide on one part of a page (e.g. just the repairs card,
    without the next section bleeding in below).
    """
    from PIL import Image
    src = os.path.join(SHOTS, name + ".png")
    if not os.path.exists(src):
        return name
    stem = f"{name}--crop{int(top * 100)}-{int(bottom * 100)}"
    dst = os.path.join(SHOTS, stem + ".png")
    im = Image.open(src)
    im.crop((0, int(im.height * top), im.width,
             int(im.height * bottom))).save(dst)
    return stem


def notes(s, text):
    s.notes_slide.notes_text_frame.text = text


def accent(s, x, y, w=Inches(1.05), h=Inches(0.07)):
    r = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    r.fill.solid()
    r.fill.fore_color.rgb = RED
    r.line.fill.background()
    r.shadow.inherit = False
    return r


def badge(s, x, y, text, w=Inches(1.5), h=Inches(0.36), fill=RED, fg=WHITE):
    r = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    r.fill.solid()
    r.fill.fore_color.rgb = fill
    r.line.fill.background()
    r.shadow.inherit = False
    tf = r.text_frame
    tf.margin_left = tf.margin_right = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.text = text
    p.alignment = PP_ALIGN.CENTER
    p.font.size, p.font.bold, p.font.color.rgb = Pt(12), True, fg
    p.font.name = "Calibri"
    return r


def footer(s, n):
    textbox(s, Inches(11.6), Inches(6.95), Inches(1.25), Inches(0.3),
            f"{n}", size=11, color=MUTED, align=PP_ALIGN.RIGHT)
    textbox(s, Inches(0.62), Inches(6.95), Inches(6), Inches(0.3),
            "Lucky Lion eBike — Web System", size=11, color=MUTED)


# ------------------------------------------------------------- slide recipes
def title_slide():
    s = add_slide(INK)
    band = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Inches(2.42), W,
                              Inches(0.1))
    band.fill.solid()
    band.fill.fore_color.rgb = RED
    band.line.fill.background()
    band.shadow.inherit = False

    textbox(s, Inches(1.1), Inches(1.35), Inches(11), Inches(0.5),
            "CAPSTONE PROJECT  ·  TRECE MARTIRES, CAVITE",
            size=14, bold=True, color=RED)
    textbox(s, Inches(1.1), Inches(2.72), Inches(11.2), Inches(1.6),
            "Lucky Lion eBike\nWeb Management System",
            size=46, bold=True, color=WHITE, line=1.12)
    textbox(s, Inches(1.1), Inches(4.55), Inches(10), Inches(1.2),
            "An online catalog, customer member portal, and staff "
            "back-office — built on one connected database.",
            size=18, color=RGBColor(0xC9, 0xCC, 0xD2), line=1.35)
    textbox(s, Inches(1.1), Inches(5.95), Inches(10), Inches(0.5),
            "Presented to Lucky Lion Main Office", size=15,
            bold=True, color=WHITE)
    notes(s, "Introduce the team and the branch. One line: this is a working "
             "system, not a mockup — everything you'll see is live data.")


def divider(n, kicker, title, sub):
    s = add_slide(INK)
    accent(s, Inches(1.1), Inches(2.55), w=Inches(1.4), h=Inches(0.09))
    textbox(s, Inches(1.1), Inches(2.0), Inches(10), Inches(0.4),
            kicker, size=14, bold=True, color=RED)
    textbox(s, Inches(1.1), Inches(2.95), Inches(11), Inches(1.0),
            title, size=38, bold=True, color=WHITE)
    textbox(s, Inches(1.1), Inches(4.15), Inches(10), Inches(0.8),
            sub, size=17, color=RGBColor(0xC9, 0xCC, 0xD2), line=1.35)
    footer(s, n)
    return s


def feature(n, title, sub, shot, points, tag=None):
    """Left copy column, right screenshot."""
    s = add_slide()
    accent(s, Inches(0.62), Inches(0.62))
    textbox(s, Inches(0.62), Inches(0.85), Inches(4.0), Inches(0.9),
            title, size=27, bold=True, line=1.1)
    textbox(s, Inches(0.62), Inches(1.95), Inches(3.5), Inches(0.8),
            sub, size=14, color=MUTED, line=1.3)
    # The badge sits above the bullets: bullet blocks vary in length, and
    # anchoring it to the bottom let long ones run into it.
    y = 2.72
    if tag:
        w = Inches(0.34 + 0.098 * len(tag))
        badge(s, Inches(0.62), Inches(y), tag, w=w)
        y += 0.72
    bullets(s, Inches(0.62), Inches(y), Inches(3.5), points, size=14.5)
    picture(s, shot, Inches(4.5), Inches(0.95), w=Inches(8.2))
    footer(s, n)
    return s


def flow(n, step, title, sub, shot, caption):
    """Booking-story slide: numbered step, big screenshot."""
    s = add_slide()
    badge(s, Inches(0.62), Inches(0.6), f"STEP {step}", w=Inches(1.25))
    textbox(s, Inches(2.05), Inches(0.58), Inches(8.5), Inches(0.5),
            title, size=25, bold=True)
    textbox(s, Inches(2.05), Inches(1.12), Inches(9.5), Inches(0.4),
            sub, size=14, color=MUTED)
    # Fit inside the band to the right of the callout column: size by height,
    # but fall back to width for wide crops so it never overflows left.
    band_x, band_w, max_h = Inches(2.75), Inches(9.9), Inches(4.55)
    pic = picture(s, shot, band_x, Inches(1.62), h=max_h)
    if pic:
        if pic.width > band_w:
            scale = band_w / pic.width
            pic.width, pic.height = band_w, int(pic.height * scale)
        pic.left = int(band_x + (band_w - pic.width) / 2)
        pic.top = int(Inches(1.62) + (max_h - pic.height) / 2)
    textbox(s, Inches(2.75), Inches(6.33), Inches(9.9), Inches(0.55),
            caption, size=12, color=MUTED, align=PP_ALIGN.CENTER, line=1.2)
    footer(s, n)
    return s


def pair(n, title, sub, left_shot, left_label, right_shot, right_label,
         caption, arrow="→"):
    """Two-sided sync slide: one action on the left, its live result on the
    right, with an arrow between — the 'what staff do / what customers see'
    story the whole deck builds toward."""
    s = add_slide()
    accent(s, Inches(0.62), Inches(0.55))
    textbox(s, Inches(0.62), Inches(0.72), Inches(12.1), Inches(0.6),
            title, size=25, bold=True)
    textbox(s, Inches(0.62), Inches(1.42), Inches(12.1), Inches(0.4),
            sub, size=14, color=MUTED)

    panel_w = Inches(5.55)
    lx, rx = Inches(0.62), Inches(7.22)
    lab_y, img_top, max_h = Inches(2.28), Inches(2.82), Inches(3.5)

    badge(s, lx, lab_y, left_label, w=panel_w, h=Inches(0.42), fill=INK)
    badge(s, rx, lab_y, right_label, w=panel_w, h=Inches(0.42), fill=RED)

    def place(shot, x):
        p = picture(s, shot, x, img_top, w=panel_w)
        if not p:
            return None
        if p.height > max_h:
            scale = max_h / p.height
            p.width, p.height = int(p.width * scale), int(max_h)
        p.left = int(x + (panel_w - p.width) / 2)
        return p

    place(left_shot, lx)
    place(right_shot, rx)

    # arrow disc in the gap between the two panels, at image mid-height
    mid = img_top + max_h / 2
    disc = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(6.28),
                              int(mid - Inches(0.44)), Inches(0.88),
                              Inches(0.88))
    disc.fill.solid()
    disc.fill.fore_color.rgb = RED
    disc.line.color.rgb = WHITE
    disc.line.width = Pt(2.5)
    disc.shadow.inherit = False
    tf = disc.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.text = arrow
    p.alignment = PP_ALIGN.CENTER
    p.font.size, p.font.bold, p.font.color.rgb = Pt(24), True, WHITE
    p.font.name = "Calibri"

    textbox(s, Inches(0.62), Inches(6.5), Inches(12.1), Inches(0.5),
            caption, size=13, color=MUTED, align=PP_ALIGN.CENTER, line=1.2)
    footer(s, n)
    return s


# --------------------------------------------------------------- build deck
def build():
    n = [1]

    def num():
        n[0] += 1
        return n[0]

    title_slide()

    # --- context -----------------------------------------------------------
    s = add_slide()
    accent(s, Inches(0.62), Inches(0.62))
    textbox(s, Inches(0.62), Inches(0.85), Inches(8), Inches(0.7),
            "The Problem & What We Built", size=27, bold=True)
    cols = [
        ("Before", [
            "Customers had to visit or message the store just to ask "
            "about models and prices.",
            "Repair requests were taken by phone or chat, with no record "
            "staff could track.",
            "Promo posters were made by hand for every sale.",
            "No single place to see customers and their purchased units.",
        ]),
        ("With the system", [
            "A public catalog with photos, specs, and all three payment "
            "options, open 24/7.",
            "Customers book repairs online; the request lands in the staff "
            "portal instantly.",
            "Staff generate a branded sale poster in seconds.",
            "One dashboard for users, units, repairs, models, and site "
            "settings.",
        ]),
    ]
    for i, (head, items) in enumerate(cols):
        x = Inches(0.62 + i * 6.3)
        card = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x,
                                  Inches(1.85), Inches(5.9), Inches(4.75))
        card.fill.solid()
        card.fill.fore_color.rgb = WHITE if i else RGBColor(0xF2, 0xF3, 0xF5)
        card.line.color.rgb = RED if i else LINE
        card.line.width = Pt(1.5 if i else 1)
        card.shadow.inherit = False
        card.text_frame.text = ""
        textbox(s, x + Inches(0.4), Inches(2.2), Inches(5), Inches(0.4),
                head.upper(), size=13, bold=True, color=RED if i else MUTED)
        bullets(s, x + Inches(0.4), Inches(2.75), Inches(5.1), items, size=13.5)
    footer(s, num())
    notes(s, "Frame the business problem first, then the fix. Keep it short — "
             "the screenshots do the convincing.")

    # --- architecture ------------------------------------------------------
    s = add_slide()
    accent(s, Inches(0.62), Inches(0.62))
    textbox(s, Inches(0.62), Inches(0.85), Inches(8), Inches(0.7),
            "How the System Is Built", size=27, bold=True)
    textbox(s, Inches(0.62), Inches(1.6), Inches(9), Inches(0.4),
            "Three layers, one shared database — which is why a customer "
            "booking shows up for staff immediately.",
            size=14, color=MUTED)
    layers = [
        ("PUBLIC WEBSITE", "index.html",
         "Catalog, About, Location, Contact — open to everyone, no login."),
        ("MEMBER PORTAL", "customer login",
         "Profile, owned units, manuals, maintenance, tips, road rules, "
         "and repair booking."),
        ("STAFF PORTAL", "staff.html",
         "Users, repairs, models & pricing, poster maker, store photos, "
         "site settings."),
        ("DATABASE & STORAGE", "Supabase (PostgreSQL)",
         "Users, e-bikes, repairs, model settings, photos, site settings — "
         "hosted online."),
    ]
    y = 2.25
    for i, (head, tech, desc) in enumerate(layers):
        last = i == len(layers) - 1
        card = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.62),
                                  Inches(y), Inches(12.1), Inches(0.95))
        card.fill.solid()
        card.fill.fore_color.rgb = RED if last else WHITE
        card.line.color.rgb = RED if last else LINE
        card.line.width = Pt(1)
        card.shadow.inherit = False
        card.text_frame.text = ""
        fg = WHITE if last else INK
        textbox(s, Inches(1.0), Inches(y + 0.16), Inches(3.1), Inches(0.35),
                head, size=13.5, bold=True, color=fg)
        textbox(s, Inches(1.0), Inches(y + 0.53), Inches(3.1), Inches(0.3),
                tech, size=11.5,
                color=RGBColor(0xFF, 0xD9, 0xDA) if last else MUTED)
        textbox(s, Inches(4.35), Inches(y + 0.3), Inches(8.1), Inches(0.5),
                desc, size=13, color=fg)
        y += 1.12
    footer(s, num())
    notes(s, "The key point for the boss: staff and customers are not looking "
             "at separate systems. One database, updated live.")

    # --- PUBLIC ------------------------------------------------------------
    divider(num(), "PART 1", "The Public Website",
            "What any customer sees when they search for Lucky Lion online.")

    feature(num(), "Home Page", "The landing page any visitor lands on.",
            "01-home-hero",
            ["Rotating photos of the actual Trece Martires store, managed by "
             "staff — no developer needed.",
             "Branch phone number, directions, and business hours right in "
             "the header.",
             "Facebook, YouTube, and Viber links, plus a floating Messenger "
             "button on every page.",
             "Seasonal sale banner that staff can switch on for promos."],
            tag="LIVE STORE PHOTOS")

    feature(num(), "E-Bike Catalog",
            "All units, grouped by wheel type.", "02-store-catalog",
            ["22 models across 2-, 3-, and 4-wheel categories.",
             "'New' and 'Sale' badges are toggled by staff in the portal.",
             "Staff control the display order, so best-sellers can be put "
             "on top.",
             "Prices shown for cash, credit card, and Home Credit."],
            tag="22 MODELS")

    feature(num(), "Model Details",
            "Tap any unit for the full spec sheet.", "03-model-viewer",
            ["Photo gallery with next/previous arrows, a photo counter, and "
             "a thumbnail strip.",
             "Full specs: wheels, motor wattage, battery, and brakes.",
             "All three payment options side by side.",
             "Photos and specs are editable by staff at any time."])

    feature(num(), "About Us",
            "The company story and its credentials.", "05-about-us-award",
            ["Company background from the Imus main branch to the Trece "
             "Martires branch.",
             "Highlights lifetime service, flexible payment, and free "
             "delivery within Cavite.",
             "Features the World Class Philippines award — Best E-Bike "
             "Marketing Company, December 2022.",
             "Builds trust with customers who are comparing sellers online."],
            tag="AWARD-WINNING")

    feature(num(), "Why Choose Lucky Lion",
            "The selling points, stated plainly.", "04-why-choose-us",
            ["Highlights the branch's advantages for first-time e-bike "
             "buyers.",
             "Content is drawn from a data file, so wording is easy to "
             "update.",
             "Sits directly above the About section to reinforce the pitch."])

    feature(num(), "Location & Branch Info",
            "Getting customers to the store.", "06-location",
            ["Complete Trece Martires branch address and business hours.",
             "One-tap link that opens Google Maps directions.",
             "Staff edit this in Site Settings — no code changes needed."])

    feature(num(), "Contact Us",
            "Every way to reach the branch.", "07-contact",
            ["Phone, email, Viber, Facebook, and YouTube in one place.",
             "Tap-to-call works directly on a phone.",
             "Messenger button floats on every page for instant chat."])

    # --- MEMBER ------------------------------------------------------------
    divider(num(), "PART 2", "The Customer Member Area",
            "What a customer gets after buying a unit and receiving an "
            "account.")

    feature(num(), "Member Login",
            "Accounts are issued by staff.", "08-login-modal",
            ["Customers do not self-register — staff create the account "
             "when a unit is purchased.",
             "Each account gets an automatic member ID (CUST-0001, and so "
             "on).",
             "'Remember me' keeps the customer signed in on their own "
             "phone.",
             "Staff accounts use the same login but open the staff portal."])

    feature(num(), "Member Menu",
            "Everything a customer owner gets.", "09-member-menu",
            ["Profile, Manuals, Maintenance, Tips, Road Rules, and Schedule "
             "Repair.",
             "Available from any page once signed in.",
             "Turns a one-time sale into an ongoing service relationship."])

    feature(num(), "Customer Profile",
            "Their account and their units.", "10-customer-profile",
            ["Shows member ID, contact details, and address.",
             "Lists every e-bike the customer owns — a customer can have "
             "more than one.",
             "Staff link each purchased unit to the account, which is what "
             "makes repair booking possible."])

    feature(num(), "Manuals & Maintenance",
            "After-sales support, self-service.", "11-manuals",
            ["Owner's manual content available any time.",
             "Maintenance schedule so customers know when servicing is due.",
             "Fewer basic 'how do I…' calls to branch staff."])

    feature(num(), "Riding Tips & Road Rules",
            "Safety and compliance guidance.", "13-tips",
            ["Practical riding and battery-care tips to extend unit life.",
             "Philippine road rules for e-bikes and e-trikes.",
             "Positions Lucky Lion as a responsible seller, not just a "
             "retailer."])

    # --- STAFF -------------------------------------------------------------
    divider(num(), "PART 3", "The Staff Portal",
            "Redesigned back-office: a Dashboard landing page and a single "
            "side menu — one tool on screen at a time, no developer required.")

    s = feature(num(), "The Staff Dashboard",
                "The branch at a glance, the moment staff sign in.",
                "17a-staff-dashboard",
                ["A time-of-day greeting, today's date, and the active "
                 "seasonal theme up top.",
                 "Four live tiles: open repairs, overdue, booked today, and "
                 "completed.",
                 "'Next repairs' lists the soonest jobs with tap-to-call; "
                 "overdue turns red.",
                 "Quick-action buttons jump straight to any task."],
                tag="NEW LANDING PAGE")
    notes(s, "This whole page is new. Everything the staff need first thing "
             "in the morning — what's due, what's overdue — is on one screen.")

    feature(num(), "One Menu, Every Tool",
            "The new grouped side drawer.", "17b-staff-drawer",
            ["Tools are grouped by job: Customers, Store, Marketing, System.",
             "The menu tucks away so staff see one focused screen at a time.",
             "A live badge shows how many repairs are still open.",
             "A 'Back to Dashboard' button sits on every screen."],
            tag="REDESIGNED NAVIGATION")

    feature(num(), "Users & Their E-Bikes",
            "The branch's customer list.", "19-staff-users",
            ["Every customer with their contact details and purchased "
             "units.",
             "Add or remove units as customers buy more.",
             "Staff members have profile photos; customers are listed by "
             "member ID.",
             "Edit details or reset a password when a customer forgets it."])

    feature(num(), "Creating Accounts",
            "Onboarding a new buyer.", "20-staff-create-user",
            ["Create a customer account at the moment of purchase.",
             "Member IDs are generated automatically (CUST-#### / "
             "STAFF-####).",
             "The purchased unit is linked right away, which enables repair "
             "booking.",
             "Staff can also create accounts for new staff members."])

    feature(num(), "Managing Repairs",
            "From request to completion.", "17-staff-repairs-scheduled",
            ["Move a job through pending → in-progress → done in one tap.",
             "Overdue bookings are highlighted in red so nothing is "
             "forgotten.",
             "Edit the repair type, date, notes, and the amount charged.",
             "Cancel, reopen, or delete a booking when plans change."],
            tag="STATUS TRACKING")

    feature(num(), "Completed Repair History",
            "A permanent service record.", "18-staff-repairs-history",
            ["Every finished repair is kept with its completion date.",
             "Searchable by customer or unit.",
             "Totals up the charges logged, giving the branch a simple "
             "record of service income.",
             "Useful for warranty checks and for spotting repeat issues."],
            tag="SEARCHABLE + TOTALS")

    feature(num(), "Models & Pricing Manager",
            "Full control of the catalog.", "21-staff-models",
            ["Edit any model's name, wheel type, motor, battery, brakes, "
             "and price.",
             "Set sale prices for cash, credit card, and Home Credit, and "
             "toggle New / Sale badges.",
             "Upload multiple photos per unit, crop them, and drag to "
             "reorder.",
             "Add brand-new models, hide discontinued ones, and reorder how "
             "they appear in the store."],
            tag="NO DEVELOPER NEEDED")

    feature(num(), "Poster Maker",
            "Marketing material in seconds.", "22-poster-form",
            ["Pick a unit and a seasonal theme; every field is pre-filled "
             "from the catalog.",
             "Each field stays editable, so staff can adjust the wording "
             "before generating.",
             "Contact details fill in automatically from Site Settings.",
             "Generates a ready-to-post image for Facebook."],
            tag="READY TO POST")

    s = add_slide()
    accent(s, Inches(0.62), Inches(0.62))
    textbox(s, Inches(0.62), Inches(0.85), Inches(9), Inches(0.7),
            "Poster Maker — The Result", size=27, bold=True)
    textbox(s, Inches(0.62), Inches(1.58), Inches(11.5), Inches(0.4),
            "Filled in on the left, generated on the right — then downloaded "
            "as an image file and posted straight to Facebook.",
            size=14, color=MUTED)
    # the form on the left, the finished artwork (captured straight from the
    # poster canvas, so nothing is cropped) on the right
    fp = picture(s, crop("22-poster-form", 0.14, 1.0), Inches(0.62),
                 Inches(2.15), w=Inches(7.2))
    art = picture(s, "23b-poster-artwork", Inches(0), Inches(2.15),
                  h=Inches(4.35))
    if art:
        art.left = int(Inches(12.71) - art.width)
    if fp and art and fp.left + fp.width > art.left - Inches(0.3):
        fp.width = int(art.left - Inches(0.3) - fp.left)
        fp.height = int(fp.width / 1.6)
    footer(s, num())
    notes(s, "Worth stressing the time saved: no graphic designer, no Canva "
             "subscription, and the prices always match the live catalog.")

    feature(num(), "Store Photos",
            "Keeping the home page fresh.", "24-staff-store-photos",
            ["Upload new store photos with a built-in cropping tool.",
             "Reorder or delete photos; the home page updates immediately.",
             "Keeps the site looking current without touching any code."])

    feature(num(), "Site Settings",
            "One place for branch information.", "25-staff-settings",
            ["Phone, email, address, business hours, and map link.",
             "Facebook, YouTube, and Viber links.",
             "Seasonal theme switch for Christmas, summer, and other promos.",
             "Change it once and it updates across the entire website."],
            tag="UPDATE ONCE, SHOWS EVERYWHERE")

    # --- ONE DATABASE, TWO SIDES (live sync) -------------------------------
    divider(num(), "PART 4", "One Database, Two Sides",
            "The heart of the system: what staff do on their screen shows up "
            "for customers right away — and the other way around. No "
            "re-typing, no separate systems.")

    s = pair(num(),
             "Book a repair → it lands in the staff portal",
             "A customer schedules a repair from their phone; the exact "
             "request appears for staff to act on.",
             # both sides left uncropped: matching ~1.6 aspect keeps the two
             # panels the same height, so the pair reads as one comparison
             "15-schedule-repair-form", "CUSTOMER BOOKS",
             "17-staff-repairs-scheduled", "STAFF SEE IT INSTANTLY",
             "The request arrives as a live job — unit, date, and time slot, "
             "with the customer's number and address ready to tap. Nothing "
             "is re-typed by hand.")
    notes(s, "The centrepiece. Offer to prove it live at the end: book a "
             "repair on a phone, refresh the portal, and the new row appears "
             "with the same unit, date and notes.")

    s = pair(num(),
             "Flip the seasonal theme → the whole site changes",
             "Staff pick a theme in Site Settings; every visitor's homepage "
             "gets the seasonal look and the sale banner — no code, no "
             "re-deploy.",
             "31-staff-settings-theme", "STAFF PICK THE THEME",
             "30-home-themed", "EVERY VISITOR SEES THE SALE",
             "One dropdown drives the banner, colours, and festive touches "
             "across the entire public website in an instant.")
    notes(s, "Great visual moment: dull dropdown on the left, festive "
             "homepage with the 'Christmas Sale' ribbon on the right.")

    s = pair(num(),
             "Edit a model → the catalog updates live",
             "Staff change a price, toggle a New/Sale badge, or reorder "
             "units; the public store reflects it immediately.",
             "21-staff-models", "STAFF EDIT A MODEL",
             "02-store-catalog", "CUSTOMERS SEE IT LIVE",
             "Specs, prices, badges, photos, and display order are all "
             "controlled from the portal — the store is never out of date.")
    notes(s, "Reinforce the theme: staff are in full control of the catalog, "
             "and customers always see the current version.")

    # --- mobile ------------------------------------------------------------
    s = add_slide()
    accent(s, Inches(0.62), Inches(0.62))
    textbox(s, Inches(0.62), Inches(0.85), Inches(9), Inches(0.7),
            "Works on Phones Too", size=27, bold=True)
    textbox(s, Inches(0.62), Inches(1.58), Inches(5.4), Inches(1.4),
            "Most Filipino customers will open this on a phone, not a "
            "laptop. Every page — including repair booking — is built to "
            "work at phone size.",
            size=14.5, color=MUTED, line=1.35)
    bullets(s, Inches(0.62), Inches(3.1), Inches(5.6),
            ["The menu collapses into a tap-to-open button.",
             "Model cards stack into a single readable column.",
             "Tap-to-call and Get Directions work natively on mobile.",
             "The full booking form is usable one-handed."], size=14)
    # Phone shots are tall. Place them at a trial height, measure the real
    # width, then scale so the row always fits the space to the right of the
    # copy — whatever aspect ratio the captures happen to have.
    phones = ["26-mobile-home", "27-mobile-store", "29-mobile-schedule"]
    left, gap, avail = Inches(6.35), Inches(0.16), Inches(6.36)
    trial = Inches(5.0)
    pics = [p for p in (picture(s, n, left, Inches(1.5), h=trial)
                        for n in phones) if p]
    if pics:
        row = sum(p.width for p in pics) + gap * (len(pics) - 1)
        scale = min(1.0, avail / row)
        x = left
        for p in pics:
            p.height = int(trial * scale)
            p.width = int(p.width * scale)
            p.left, p.top = x, Inches(1.6)
            x += p.width + gap
    footer(s, num())
    notes(s, "Quick slide. The point is simply that nothing breaks on a "
             "phone — no need to dwell.")

    # --- status / roadmap --------------------------------------------------
    s = add_slide()
    accent(s, Inches(0.62), Inches(0.62))
    textbox(s, Inches(0.62), Inches(0.85), Inches(9), Inches(0.7),
            "Current Status & Next Steps", size=27, bold=True)
    groups = [
        ("WORKING NOW", RED, [
            "Public catalog with 22 models, live online.",
            "Customer accounts, owned-unit tracking, and repair booking.",
            "Full staff portal: users, repairs, models, poster, settings.",
            "Photo uploads and seasonal themes, all staff-controlled.",
        ]),
        ("BEING IMPROVED", MUTED, [
            "Login security is at prototype level — it needs proper "
            "password protection before wide public use.",
            "Poster Maker still uses the built-in model photos rather than "
            "staff-uploaded ones.",
            "Updates are deployed manually rather than automatically.",
        ]),
        ("PROPOSED NEXT", MUTED, [
            "SMS or email reminders for upcoming repair appointments.",
            "Sales and repair-income reporting for the branch.",
            "Online reservation or downpayment for a unit.",
            "Roll the system out to the other Lucky Lion branches.",
        ]),
    ]
    for i, (head, col, items) in enumerate(groups):
        x = Inches(0.62 + i * 4.18)
        card = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x,
                                  Inches(1.75), Inches(3.9), Inches(4.9))
        card.fill.solid()
        card.fill.fore_color.rgb = WHITE
        card.line.color.rgb = col if i == 0 else LINE
        card.line.width = Pt(1.5 if i == 0 else 1)
        card.shadow.inherit = False
        card.text_frame.text = ""
        textbox(s, x + Inches(0.35), Inches(2.1), Inches(3.3), Inches(0.4),
                head, size=12.5, bold=True, color=col)
        bullets(s, x + Inches(0.35), Inches(2.62), Inches(3.25), items,
                size=12.5)
    footer(s, num())
    notes(s, "Be honest about the prototype-level login — the boss will "
             "respect knowing the limits. Then pivot to the roll-out idea.")

    # --- close -------------------------------------------------------------
    s = add_slide(INK)
    accent(s, Inches(1.1), Inches(2.35), w=Inches(1.4), h=Inches(0.09))
    textbox(s, Inches(1.1), Inches(2.75), Inches(11), Inches(1.0),
            "Thank You", size=42, bold=True, color=WHITE)
    textbox(s, Inches(1.1), Inches(3.95), Inches(10.5), Inches(1.4),
            "One website for customers, one portal for staff, one database "
            "keeping both in sync — ready to demo live.",
            size=18, color=RGBColor(0xC9, 0xCC, 0xD2), line=1.35)
    textbox(s, Inches(1.1), Inches(5.35), Inches(10), Inches(0.9),
            "Demo accounts —  customer: user / user123      "
            "staff: staff / staff123",
            size=14, bold=True, color=RED)
    footer(s, num())
    notes(s, "Offer to run the live demo: book a repair on the phone, then "
             "show it appearing on the staff screen.")

    prs.save(OUTFILE)
    print(f"Saved {OUTFILE}  ({len(prs.slides.__iter__.__self__._sldIdLst)} "
          f"slides)")


if __name__ == "__main__":
    build()
