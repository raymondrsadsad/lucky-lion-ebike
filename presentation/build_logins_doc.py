#!/usr/bin/env python3
"""Build Lucky-Lion-Logins.docx — the demo account credentials sheet.

  python3 presentation/build_logins_doc.py

Accounts here MUST match supabase/demo-data.sql (single source of truth).
Password rule for the branch accounts: username + "123".
"""
import os
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

HERE = os.path.dirname(os.path.abspath(__file__))
OUTFILE = os.path.join(HERE, "Lucky-Lion-Logins.docx")

RED = RGBColor(0xE1, 0x1B, 0x22)
INK = RGBColor(0x14, 0x14, 0x14)
MUTED = RGBColor(0x6B, 0x70, 0x7B)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

# Branch → (customer, staff). Passwords = username + "123".
BRANCHES = [
    ("Main — Buhay na Tubig, Imus", "MAIN",   ("Maria Santos", "MAIN-C-0001", "maria"),  ("Rico Cruz",        "MAIN-S-01",  "rico")),
    ("Imus",                        "IMUS",   ("Pedro Reyes", "IMUS-C-0001", "pedro"),   ("Juan Delos Reyes", "IMUS-S-01",  "juan")),
    ("Trece Martires",              "TRECE",  ("Ana Lim",     "TRECE-C-0001", "ana"),    ("Lito Ramos",       "TRECE-S-01", "lito")),
    ("Tanza",                       "TANZA",  ("Carlo Mendoza","TANZA-C-0001", "carlo"), ("Bea Villanueva",   "TANZA-S-01", "bea")),
    ("Naic",                        "NAIC",   ("Rosa Fernandez","NAIC-C-0001", "rosa"),  ("Nico Aquino",      "NAIC-S-01",  "nico")),
    ("Tejero (General Trias)",      "TEJERO", ("Dennis Garcia","TEJERO-C-0001", "dennis"),("Grace Bautista",  "TEJERO-S-01","grace")),
    ("Gen. Mariano Alvarez (GMA)",  "GMA",    ("Mila Torres", "GMA-C-0001", "mila"),      ("Paolo Navarro",    "GMA-S-01",   "paolo")),
]

CUST_FILL = "FDEBEC"   # light red tint for customer rows
STAFF_FILL = "EEF1F6"  # light grey-blue for staff rows
HEAD_FILL = "E11B22"   # brand red header
ADMIN_FILL = "141414"  # dark for the admin highlight


def _set_font(run, size=11, bold=False, color=INK, name="Calibri"):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = name


def _shade(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


def _cell(cell, text, size=11, bold=False, color=INK, fill=None, align=None):
    cell.text = ""
    p = cell.paragraphs[0]
    if align is not None:
        p.alignment = align
    r = p.add_run(text)
    _set_font(r, size=size, bold=bold, color=color)
    if fill:
        _shade(cell, fill)


def build():
    doc = Document()

    # Page margins a touch tighter so the table breathes.
    for sec in doc.sections:
        sec.left_margin = sec.right_margin = Inches(0.7)
        sec.top_margin = sec.bottom_margin = Inches(0.7)

    # ---- Title block ----
    kicker = doc.add_paragraph()
    kr = kicker.add_run("LUCKY LION eBIKE  ·  DEMO ACCOUNTS")
    _set_font(kr, size=11, bold=True, color=RED)

    title = doc.add_paragraph()
    tr = title.add_run("Login & Password Reference")
    _set_font(tr, size=22, bold=True, color=INK)
    title.space_after = Pt(2)

    note = doc.add_paragraph()
    nr = note.add_run("Demo accounts for the walkthrough only. One customer and one staff "
                      "account per branch, plus the company admin. Passwords follow the "
                      "pattern: username + “123”.")
    _set_font(nr, size=10.5, color=MUTED)
    note.space_after = Pt(10)

    # ---- Admin callout ----
    admin = doc.add_paragraph()
    a1 = admin.add_run("ADMIN (all branches):  ")
    _set_font(a1, size=12, bold=True, color=RED)
    a2 = admin.add_run("username ")
    _set_font(a2, size=12, color=INK)
    a3 = admin.add_run("admin")
    _set_font(a3, size=12, bold=True, color=INK)
    a4 = admin.add_run("   password ")
    _set_font(a4, size=12, color=INK)
    a5 = admin.add_run("admin123")
    _set_font(a5, size=12, bold=True, color=INK)
    a6 = admin.add_run("   (Member ID ADMIN-01 · staff.html → Dashboard + Analytics)")
    _set_font(a6, size=10.5, color=MUTED)
    admin.space_after = Pt(12)

    # ---- Table ----
    headers = ["Branch", "Role", "Name", "Member ID", "Username", "Password"]
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = True

    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        _cell(hdr[i], h, size=11, bold=True, color=WHITE, fill=HEAD_FILL)

    for label, code, cust, staff in BRANCHES:
        cname, cid, cuser = cust
        sname, sid, suser = staff
        # customer row
        row = table.add_row().cells
        _cell(row[0], label, bold=True, fill=CUST_FILL)
        _cell(row[1], "Customer", fill=CUST_FILL)
        _cell(row[2], cname, fill=CUST_FILL)
        _cell(row[3], cid, fill=CUST_FILL)
        _cell(row[4], cuser, bold=True, fill=CUST_FILL)
        _cell(row[5], cuser + "123", bold=True, fill=CUST_FILL)
        # staff row
        row = table.add_row().cells
        _cell(row[0], "", fill=STAFF_FILL)
        _cell(row[1], "Staff", fill=STAFF_FILL)
        _cell(row[2], sname, fill=STAFF_FILL)
        _cell(row[3], sid, fill=STAFF_FILL)
        _cell(row[4], suser, bold=True, fill=STAFF_FILL)
        _cell(row[5], suser + "123", bold=True, fill=STAFF_FILL)

    # admin row (dark highlight)
    row = table.add_row().cells
    _cell(row[0], "All branches", bold=True, color=WHITE, fill=ADMIN_FILL)
    _cell(row[1], "Admin", color=WHITE, fill=ADMIN_FILL)
    _cell(row[2], "Owner / Admin", color=WHITE, fill=ADMIN_FILL)
    _cell(row[3], "ADMIN-01", color=WHITE, fill=ADMIN_FILL)
    _cell(row[4], "admin", bold=True, color=WHITE, fill=ADMIN_FILL)
    _cell(row[5], "admin123", bold=True, color=WHITE, fill=ADMIN_FILL)

    # ---- How to log in ----
    doc.add_paragraph().space_after = Pt(2)
    how = doc.add_paragraph()
    hr = how.add_run("How to log in")
    _set_font(hr, size=13, bold=True, color=INK)

    for line in [
        "Customers log in on the public site (index.html) — Login button, top right.",
        "Staff and Admin use the portal (staff.html); it redirects non-staff back to the public site.",
        "Admin sees every branch (read-only oversight) plus the Dashboard and Analytics.",
        "Staff see and manage only their own branch.",
    ]:
        p = doc.add_paragraph(style="List Bullet")
        pr = p.add_run(line)
        _set_font(pr, size=10.5, color=INK)

    tip = doc.add_paragraph()
    tr2 = tip.add_run("These accounts are created by supabase/demo-data.sql and removed by "
                      "supabase/demo-data-cleanup.sql after the demo.")
    _set_font(tr2, size=9.5, color=MUTED)

    doc.save(OUTFILE)
    print("wrote", OUTFILE)


if __name__ == "__main__":
    build()
